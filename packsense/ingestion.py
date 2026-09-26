"""Stops 1-2: read scenario tables without inventing missing input values.

This module checks scenario facts only. A schema-valid row is not yet a safe
recommendation: food classification and reference evidence join in later stops.
"""

import csv
import hashlib
import json
from collections import Counter
from dataclasses import dataclass
from math import isfinite
from pathlib import Path
from typing import Any, Mapping
from zipfile import BadZipFile

from openpyxl import load_workbook

from packsense.contracts import HandlingSeverity, ScenarioInput, StorageType
from packsense.units import (
    PACK_QUANTITY_UNITS,
    SCENARIO_REFERENCE_COLUMNS,
    SCENARIO_REQUIRED_COLUMNS,
    SCENARIO_RESPIRATION_COLUMNS,
)


RESPIRATION_UNITS = frozenset({"mg CO2/kg/h", "mg O2/kg/h"})
PERCENT_FIELDS = frozenset(
    {"moisture_content_pct", "oil_fat_content_pct", "storage_relative_humidity_pct"}
)
NUMERIC_FIELDS = frozenset(
    {
        "moisture_content_pct",
        "oil_fat_content_pct",
        "pH",
        "desired_shelf_life_days",
        "storage_temperature_c",
        "storage_relative_humidity_pct",
        "transport_duration_hours",
        "transport_temperature_c",
        "transport_max_temperature_c",
        "net_pack_quantity",
        "respiration_rate",
        "respiration_reference_temperature_c",
    }
)


class InputSchemaError(ValueError):
    """A file or header problem prevents reliable row interpretation."""


@dataclass(frozen=True, slots=True)
class IngestionIssue:
    row_number: int
    record_id: str | None
    field: str
    code: str
    message: str


@dataclass(frozen=True, slots=True)
class ParsedScenarioRow:
    row_number: int
    raw_values: Mapping[str, Any]
    scenario: ScenarioInput | None
    issues: tuple[IngestionIssue, ...]


@dataclass(frozen=True, slots=True)
class ScenarioAudit:
    source_path: str
    source_sha256: str
    sheet_name: str | None
    ignored_columns: tuple[str, ...]
    rows: tuple[ParsedScenarioRow, ...]

    def report(self) -> dict[str, Any]:
        """JSON-safe exception report; no raw food values are exported."""
        issues = [issue for row in self.rows for issue in row.issues]
        counts = Counter((issue.field, issue.code) for issue in issues)
        return {
            "source_path": self.source_path,
            "source_sha256": self.source_sha256,
            "sheet_name": self.sheet_name,
            "total_rows": len(self.rows),
            "schema_valid_rows": sum(row.scenario is not None for row in self.rows),
            "rejected_rows": sum(row.scenario is None for row in self.rows),
            "ignored_columns": list(self.ignored_columns),
            "issue_counts": [
                {"field": field, "code": code, "count": count}
                for (field, code), count in sorted(counts.items())
            ],
            "issues": [
                {
                    "row_number": issue.row_number,
                    "record_id": issue.record_id,
                    "field": issue.field,
                    "code": issue.code,
                    "message": issue.message,
                }
                for issue in issues
            ],
        }


@dataclass(frozen=True, slots=True)
class _TableRow:
    number: int
    values: tuple[Any, ...]
    formula_columns: frozenset[int] = frozenset()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> tuple[tuple[Any, ...], list[_TableRow], None]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        try:
            header = tuple(next(reader))
        except StopIteration as exc:
            raise InputSchemaError("CSV is empty") from exc
        rows = [
            _TableRow(number, tuple(values))
            for number, values in enumerate(reader, start=2)
        ]
    return header, rows, None


def _read_xlsx(
    path: Path, sheet_name: str | None
) -> tuple[tuple[Any, ...], list[_TableRow], str]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise InputSchemaError("select a sheet explicitly for a multi-sheet workbook")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise InputSchemaError(f"sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows()
        try:
            header = tuple(cell.value for cell in next(iterator))
        except StopIteration as exc:
            raise InputSchemaError("worksheet is empty") from exc
        rows = []
        for number, cells in enumerate(iterator, start=2):
            rows.append(
                _TableRow(
                    number,
                    tuple(cell.value for cell in cells),
                    frozenset(
                        index for index, cell in enumerate(cells) if cell.data_type == "f"
                    ),
                )
            )
    finally:
        workbook.close()
    return header, rows, sheet_name


def _header(raw: tuple[Any, ...]) -> tuple[tuple[str, ...], tuple[str, ...]]:
    if any(not isinstance(value, str) or not value.strip() for value in raw):
        raise InputSchemaError("header names must be non-empty text")
    names = tuple(value.strip() for value in raw)
    duplicates = sorted(name for name, count in Counter(names).items() if count > 1)
    if duplicates:
        raise InputSchemaError(f"duplicate header names: {', '.join(duplicates)}")
    missing = sorted(set(SCENARIO_REQUIRED_COLUMNS) - set(names))
    if missing:
        raise InputSchemaError(f"missing required columns: {', '.join(missing)}")
    respiration = set(SCENARIO_RESPIRATION_COLUMNS) & set(names)
    if respiration and respiration != set(SCENARIO_RESPIRATION_COLUMNS):
        raise InputSchemaError("respiration columns must appear together")
    ignored = tuple(
        name
        for name in names
        if name not in (
            SCENARIO_REQUIRED_COLUMNS
            + SCENARIO_RESPIRATION_COLUMNS
            + SCENARIO_REFERENCE_COLUMNS
        )
    )
    return names, ignored


def _missing(value: Any) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _number(value: Any, field: str) -> float:
    if isinstance(value, bool):
        raise ValueError("must be a finite number")
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError("must be a finite number") from exc
    if not isfinite(number):
        raise ValueError("must be a finite number")
    if field in PERCENT_FIELDS and not 0 <= number <= 100:
        raise ValueError("must be between 0 and 100 percent")
    if field == "pH" and not 0 <= number <= 14:
        raise ValueError("must be between 0 and 14")
    if field in {"desired_shelf_life_days", "net_pack_quantity", "respiration_rate"} and number <= 0:
        raise ValueError("must be positive")
    if field == "transport_duration_hours" and number < 0:
        raise ValueError("must not be negative")
    if field.endswith("temperature_c") and number < -273.15:
        raise ValueError("cannot be below absolute zero")
    return number


def _canonical_quantity(amount: float, unit: str) -> tuple[float, str]:
    """Keep mass and volume distinct while normalizing kg/L to g/mL."""
    if unit not in PACK_QUANTITY_UNITS:
        raise ValueError("use g, kg, mL, or L")
    converted = amount * 1000 if unit in {"kg", "L"} else amount
    if not isfinite(converted):
        raise ValueError("converted quantity is not finite")
    return converted, {"kg": "g", "L": "mL"}.get(unit, unit)


def _parse_row(
    raw: Mapping[str, Any], row_number: int, duplicate: bool, formula_fields: set[str]
) -> ParsedScenarioRow:
    issues: list[IngestionIssue] = []
    record_id = raw.get("record_id")
    label = record_id.strip() if isinstance(record_id, str) else None

    def issue(field: str, code: str, message: str) -> None:
        issues.append(IngestionIssue(row_number, label, field, code, message))

    if duplicate:
        issue("record_id", "duplicate_id", "record_id appears more than once")
    for field in sorted(formula_fields):
        issue(field, "formula_cell", "formulas are not accepted as source values")
    parsed: dict[str, Any] = {}
    for field in SCENARIO_REQUIRED_COLUMNS + SCENARIO_RESPIRATION_COLUMNS + SCENARIO_REFERENCE_COLUMNS:
        value = raw.get(field)
        if _missing(value):
            if field in SCENARIO_REQUIRED_COLUMNS:
                issue(field, "missing_value", "required scenario value is blank")
            continue
        if field in formula_fields:
            continue
        if field in NUMERIC_FIELDS:
            try:
                parsed[field] = _number(value, field)
            except ValueError as exc:
                issue(field, "invalid_number", str(exc))
        elif not isinstance(value, str):
            issue(field, "invalid_text", "must be text")
        else:
            parsed[field] = value.strip()

    if "storage_type" in parsed:
        try:
            parsed["storage_type"] = StorageType(parsed["storage_type"].lower())
        except ValueError:
            issue("storage_type", "invalid_category", "use ambient, chilled, or frozen")
    if "transport_handling_severity" in parsed:
        try:
            parsed["transport_handling_severity"] = HandlingSeverity(
                parsed["transport_handling_severity"].lower()
            )
        except ValueError:
            issue("transport_handling_severity", "invalid_category", "use low, medium, or high")
    if "net_pack_quantity_unit" in parsed:
        unit = parsed["net_pack_quantity_unit"]
        if unit not in PACK_QUANTITY_UNITS:
            issue("net_pack_quantity_unit", "unsupported_unit", "use g, kg, mL, or L")
        elif "net_pack_quantity" in parsed:
            try:
                parsed["net_pack_quantity"], parsed["net_pack_quantity_unit"] = (
                    _canonical_quantity(parsed["net_pack_quantity"], unit)
                )
            except ValueError as exc:
                issue("net_pack_quantity", "invalid_number", str(exc))

    respiration_present = [not _missing(raw.get(field)) for field in SCENARIO_RESPIRATION_COLUMNS]
    if any(respiration_present) and not all(respiration_present):
        for field, present in zip(SCENARIO_RESPIRATION_COLUMNS, respiration_present):
            if not present:
                issue(field, "incomplete_respiration", "rate, unit, and reference temperature travel together")
    if "respiration_rate_unit" in parsed and parsed["respiration_rate_unit"] not in RESPIRATION_UNITS:
        issue("respiration_rate_unit", "unsupported_unit", "declare mg CO2/kg/h or mg O2/kg/h")
    if (
        "transport_temperature_c" in parsed
        and "transport_max_temperature_c" in parsed
        and parsed["transport_max_temperature_c"] < parsed["transport_temperature_c"]
    ):
        issue("transport_max_temperature_c", "inconsistent_temperature", "maximum is below transport mean")
    if parsed.get("storage_type") is StorageType.FROZEN and parsed.get("storage_temperature_c", 0) > 0:
        issue("storage_temperature_c", "inconsistent_temperature", "frozen storage is above 0 C")

    scenario = None
    if not issues:
        try:
            scenario = ScenarioInput(**parsed)
        except (TypeError, ValueError) as exc:
            issue("row", "contract_violation", str(exc))
    return ParsedScenarioRow(row_number, dict(raw), scenario, tuple(issues))


def audit_scenarios(path: str | Path, *, sheet_name: str | None = None) -> ScenarioAudit:
    """Read a CSV/XLSX scenario table and keep every original row and exception.

    Existing food reference sheets are deliberately not auto-completed into
    recommendation scenarios. `schema_valid_rows` means syntax only; produce
    classification and source joins must still run before recommendations.
    """
    source = Path(path).resolve(strict=True)
    if source.suffix.lower() == ".csv":
        if sheet_name is not None:
            raise InputSchemaError("sheet_name is only valid for XLSX")
        raw_header, table_rows, selected_sheet = _read_csv(source)
    elif source.suffix.lower() == ".xlsx":
        raw_header, table_rows, selected_sheet = _read_xlsx(source, sheet_name)
    else:
        raise InputSchemaError("input must be a .csv or .xlsx file")
    header, ignored = _header(raw_header)
    seen_ids: set[str] = set()
    parsed_rows = []
    for row in table_rows:
        if all(_missing(value) for value in row.values):
            continue
        raw = dict(zip(header, row.values))
        formula_fields = {header[index] for index in row.formula_columns if index < len(header)}
        identifier = raw.get("record_id")
        key = identifier.strip() if isinstance(identifier, str) else ""
        duplicate = bool(key and key in seen_ids)
        if key:
            seen_ids.add(key)
        parsed = _parse_row(raw, row.number, duplicate, formula_fields)
        if len(row.values) != len(header):
            length_issue = IngestionIssue(
                row.number, key or None, "row", "column_count", "row length differs from header"
            )
            parsed = ParsedScenarioRow(
                row.number, raw, None, parsed.issues + (length_issue,)
            )
        parsed_rows.append(parsed)
    return ScenarioAudit(str(source), _sha256(source), selected_sheet, ignored, tuple(parsed_rows))


def main() -> int:
    """Run an audit without changing the supplied workbook."""
    import argparse

    parser = argparse.ArgumentParser(description="Audit PackSense scenario CSV/XLSX input")
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet", help="XLSX sheet name; required if there are several sheets")
    parser.add_argument("--report", type=Path, help="write full JSON exception report")
    args = parser.parse_args()
    try:
        audit = audit_scenarios(args.input, sheet_name=args.sheet)
    except (InputSchemaError, OSError, csv.Error, BadZipFile) as exc:
        parser.exit(2, f"input error: {exc}\n")
    report = audit.report()
    if args.report:
        try:
            with args.report.open("x", encoding="utf-8") as output:
                output.write(json.dumps(report, indent=2) + "\n")
        except FileExistsError:
            parser.exit(2, "report error: output already exists; choose a new path\n")
    print(json.dumps({key: value for key, value in report.items() if key != "issues"}, indent=2))
    return 0 if not report["rejected_rows"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
